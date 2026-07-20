"""
EURES beta - Flask API Server

Dedicated deployment for the EURES beta forms and public statistics.

Environment variables:
  GRIST_API_KEY - Default API key (or GRIST_API_KEY_EURES_BETA)
  GRIST_DOC_EURES_BETA - Grist document ID used by EURES beta
  GRIST_TABLE_EURES_BETA_CANDIDATE - Candidate table (defaults to Reponses if omitted)
  GRIST_TABLE_EURES_BETA_EMPLOYER - Employer table (defaults to Reponses if omitted)
  GRIST_TABLE_EURES_BETA_STATS - Optional monthly stats table
  ADMIN_USERNAME_EURES_BETA - Admin username for /admin/eures-beta/
  ADMIN_PASSWORD_EURES_BETA - Admin password for /admin/eures-beta/
  GRIST_BASE_URL - Grist instance URL (defaults to grist.numerique.gouv.fr)
"""

import os
import json
import time
from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO
from functools import wraps
from pathlib import Path
from urllib.parse import urljoin
from dotenv import load_dotenv
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
import requests
from flask import Flask, request, jsonify, redirect, send_file, send_from_directory, Response, session, url_for, render_template_string
from markupsafe import escape
from werkzeug.middleware.dispatcher import DispatcherMiddleware

load_dotenv()

BASE_DIR = Path(__file__).parent
FORMS_DIR = BASE_DIR / 'forms'
ASSETS_DIR = BASE_DIR / 'assets'
DOCS_DIR = BASE_DIR / 'docs'

app = Flask(__name__)
app.secret_key = (
    os.environ.get('SESSION_SECRET')
    or os.environ.get('FLASK_SECRET_KEY')
    or 'dev-session-secret'
)
app.config['SESSION_COOKIE_SECURE'] = os.environ.get('SESSION_COOKIE_SECURE', '').strip().lower() in {'1', 'true', 'yes', 'oui'}
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'

GRIST_BASE_URL = os.environ.get('GRIST_BASE_URL', 'https://grist.numerique.gouv.fr').rstrip('/')
APP_MODE = os.environ.get('APP_MODE', 'eures-beta').strip().lower() or 'eures-beta'
_TABLE_COLUMNS_CACHE: dict[tuple[str, str], dict[str, object]] = {}
_TABLE_COLUMNS_CACHE_TTL_SECONDS = 60
EURES_CANDIDATS_TABLE = 'Candidats'
EURES_BESOINS_TABLE = 'Besoins_Employeurs'
EURES_MATCHINGS_TABLE = 'Matchings'
EURES_STATS_TABLE_DEFAULT = 'Pilotage_EURES_Mensuel'
EURES_MATCHING_FIELDS = {
    'besoin_id',
    'candidat_id',
    'score',
    'score_metier',
    'score_langues',
    'score_mobilite',
    'score_disponibilite',
    'score_salaire',
    'statut',
    'raisons',
    'points_faibles',
    'date_calcul',
}
EURES_SECTOR_CANONICAL_MAP = {
    'vente': 'vente',
    'commerce': 'vente',
    'sales': 'vente',
    'retail': 'vente',
    'nettoyage': 'nettoyage',
    'entretien': 'nettoyage',
    'cleaning': 'nettoyage',
    'maintenance': 'nettoyage',
    'hôtellerie': 'hotellerie',
    'hotellerie': 'hotellerie',
    'restauration': 'hotellerie',
    'hospitality': 'hotellerie',
    'catering': 'hotellerie',
    'agriculture': 'agriculture',
    'récolte': 'agriculture',
    'recolte': 'agriculture',
    'harvesting': 'agriculture',
    'polyvalent': 'polyvalent',
    'multi': 'polyvalent',
    'accessible rapidement': 'polyvalent',
}
EURES_CANDIDAT_SECTOR_SALARY_FIELDS = {
    'vente': ('tally_q20_salary_type', 'tally_q20_salary_min'),
    'nettoyage': ('tally_q22_salary_type', 'tally_q22_salary_min'),
    'hotellerie': ('tally_q25_salary_type', 'tally_q25_salary_min'),
    'agriculture': ('tally_q27_salary_type', 'tally_q27_salary_min'),
    'polyvalent': ('tally_q29_salary_type', 'tally_q29_salary_min'),
}
EURES_EMPLOYEUR_SECTOR_SALARY_FIELDS = {
    'vente': ('tally_q10_salary_type', 'tally_q10_salary_min', 'tally_q10_salary_max'),
    'nettoyage': ('tally_q11_salary_type', 'tally_q11_salary_min', 'tally_q11_salary_max'),
    'hotellerie': ('tally_q12_salary_type', 'tally_q12_salary_min', 'tally_q12_salary_max'),
    'agriculture': ('tally_q13_salary_type', 'tally_q13_salary_min', 'tally_q13_salary_max'),
    'polyvalent': ('tally_q14_salary_type', 'tally_q14_salary_min', 'tally_q14_salary_max'),
}
WIZARD_STATE_KEY = '__wizard_v3_state'
JSON_EXPORT_COLUMNS = {
    'metiers_json',
    'prestations_orp_json',
    'prestations_indirectes_json',
    'prestations_json',
    'catalogue_formations_json',
    'emploi_indicateurs_json',
    'prestations_details_json',
    'finess_json',
}
FIELD_LABELS = {
    'uuid': 'Identifiant de reprise',
    'es_nom': "Nom de l'établissement",
    'es_departement': 'Département',
    'finess_main': 'FINESS principal',
    'validateur_nom': 'Nom du validateur',
    'validateur_prenom': 'Prénom du validateur',
    'validateur_email': 'Email du validateur',
    'saisie_terminee': 'Saisie terminée',
    'check_esrp': 'Dispositif ESRP',
    'check_espo': 'Dispositif ESPO',
    'check_ueros': 'Dispositif UEROS',
    'check_deac': 'Dispositif DEAc',
    'etp_esrp': 'ETP ESRP',
    'etp_espo': 'ETP ESPO',
    'etp_ueros': 'ETP UEROS',
    'etp_deac': 'ETP DEAc',
    'q32_implantation': "Contexte social - Zone d'implantation",
    'q33_transports': 'Contexte social - Transports',
    'q33_pmr': 'Contexte social - Accessibilité PMR',
    'q33_alternatif': 'Contexte social - Transport alternatif',
    'q34_prefecture': 'Contexte social - Préfecture',
    'q35_hebergement': 'Contexte social - Hébergement',
    'q35_places': 'Contexte social - Nombre de places',
    'q35_weekend': 'Contexte social - Hébergement le week-end',
    'q36_restaurant': 'Contexte social - Restauration',
    'q37_cuisine': 'Contexte social - Cuisine',
    'q38_dui': 'Contexte social - DUI',
    'q40_remuneration': 'Contexte social - Rémunération',
    'q40_operateur': 'Contexte social - Opérateur rémunération',
    'autre_dispositif_eval': "Autre dispositif d'évaluation",
    'autre_dispositif_eval_avec_orp_cdaph': "Autre dispositif d'évaluation avec ORP CDAPH",
    'autre_dispositif_eval_sans_orp_cdaph': "Autre dispositif d'évaluation sans ORP CDAPH",
}
COH_LABELS = {
    'genre': ('Genre', ['Hommes', 'Femmes', 'Autre']),
    'age': ('Âge', ['16-17 ans', '18-19 ans', '20-24 ans', '25-29 ans', '30-34 ans', '35-39 ans', '40-44 ans', '45-49 ans', '50-54 ans', '55-59 ans', '60 ans et +', 'Je ne sais pas']),
    'niveau_entree': ("Niveau de formation à l'entrée", ['Niveau 2 (CEP, sans formation)', 'Niveau 3 (CAP, BEP)', 'Niveau 4 (Bac, Bac pro)', 'Niveau 5 (Bac +2, BTS)', 'Niveau 6 (Bac +3, Bac +4)', 'Niveau 7 (Bac +5)', 'Je ne sais pas']),
    'situation_entree': ("Situation à l'entrée", ['En emploi du secteur privé', 'Dont entreprises adaptées', 'Dont alternance (secteur privé)', 'En emploi secteur public', 'FPT', 'Dont contractuels (FPT)', 'Dont alternance (FPT)', 'FPH', 'Dont contractuels (FPH)', 'Dont alternance (FPH)', 'FPE', 'Dont contractuels (FPE)', 'Dont alternance (FPE)', 'En activité non salariée', "Travailleurs d'ESAT", "Sans emploi depuis moins d'1 an", 'Sans emploi depuis 1 à 2 ans', 'Sans emploi depuis 2 ans ou plus', "N'avaient jamais travaillé", 'En formation', 'Je ne sais pas']),
    'ressources_entree': ("Ressources à l'entrée", ['Salaire', 'AAH', 'Allocations chômage', 'RSA', 'Indemnités Journalières', 'Rentes AT/MP / Pension invalidité', 'Autres', 'Aucune', 'Je ne sais pas']),
    'pathologies': ('Pathologies', ['1 pathologie', '2 pathologies', '3 pathologies et plus', 'Non connu']),
    'origine_handicap': ('Origine du handicap', ['Congénitale', 'Maladie', 'Accident de vie privée', 'Accident travail / Pro', 'Maladie professionnelle', 'Autres', 'Non connue']),
    'lesion_origine': ('Origine de la lésion cérébrale (UEROS)', ['Traumatisme crânien', 'AVC', 'Tumeur cérébrale', 'Epilepsie', 'Autres pathologies neuro', 'Non connue']),
}
HANDICAP_TYPES = [
    'Déficiences intellectuelles',
    'Autisme et autres TED',
    'Troubles psychiques',
    'Troubles langage/apprentissages',
    'Déficiences auditives',
    'Déficiences visuelles',
    'Déficiences motrices',
    'Déficiences métaboliques/nutritionnelles',
    'Cérébro-lésions',
    'Polyhandicap',
    'TCC (comportement/com.)',
    'Diagnostics en cours',
    'Autres types de déficiences',
    'Je ne sais pas',
]
EXPORT_KEY_LABELS = {
    'done': 'Statut complété',
    'fileActive': 'File active',
    'preaccueilSansSuite': 'Pré-accueils sans suite',
    'sortiesAvantTerme': 'Sorties définitives avant le terme de la prestation',
    'sortiesTerme': 'Sorties au terme de la prestation',
    'sorties': 'Sorties définitives en 2025',
    'journees': 'Journées réalisées',
    'journeesTheoriques': 'Journées théoriques',
    'enabled': 'Activé',
    'beneficiaires': 'Nombre de personnes bénéficiaires',
    'discontinue_personnes': 'Nombre de personnes accompagnées en discontinue',
    'presentiel_total': 'Présentiel - total',
    'presentiel_complet': 'Présentiel - temps complet',
    'presentiel_partiel': 'Présentiel - temps partiel',
    'hybride_total': 'Hybride - total',
    'hybride_complet': 'Hybride - temps complet',
    'hybride_partiel': 'Hybride - temps partiel',
    'distanciel_total': 'Distanciel - total',
    'distanciel_complet': 'Distanciel - temps complet',
    'distanciel_partiel': 'Distanciel - temps partiel',
    'hors_murs_personnes': 'Hors les murs - personnes accompagnées',
    'hors_murs_journees': 'Hors les murs - journées',
    'hebergees_personnes': 'Hébergement - personnes hébergées',
    'hebergees_journees': 'Hébergement - journées',
    'hebergees_nuitees': 'Hébergement - nuitées',
    'activites_intermediaires_site': 'Activité intermédiaire - sur site',
    'activites_intermediaires_ambulatoire': 'Activité intermédiaire - en ambulatoire',
    'activites_sortie_parcours_site': 'Suite de parcours - sur site',
    'activites_sortie_parcours_ambulatoire': 'Suite de parcours - en ambulatoire',
    'droit_commun': 'Droit commun',
    'sante_social': 'Santé / Social',
    'readaptation_professionnelle': 'Réadaptation professionnelle',
    'readaptation_ueros': 'Réadaptation professionnelle - UEROS',
    'readaptation_espo': 'Réadaptation professionnelle - ESPO',
    'readaptation_esrp': 'Réadaptation professionnelle - ESRP',
    'readaptation_dfa': 'Réadaptation professionnelle - DFA',
    'inconnu': 'Situations inconnues',
    'autre': 'Autre préconisation',
    'autres': 'Autres',
    'autres_precision': 'Précision',
    'je_ne_sais_pas': 'Je ne sais pas',
    'dept': 'Département',
    'count': 'Effectif',
    'principal': 'Principal',
    'associe': 'Associé',
}

# Public mount for the daily capture tool.
try:
    from tools.fagerh_suivi_local.app import app as fagerh_suivi_app
except Exception:
    fagerh_suivi_app = None


def is_eures_beta_only_mode() -> bool:
    """Return True when this deployment must only expose EURES beta."""
    return APP_MODE in {'eures-beta', 'eures_beta', 'eures-only', 'eures_only'}


def is_form_enabled(form_id: str) -> bool:
    """Restrict exposed forms for single-purpose deployments."""
    if is_eures_beta_only_mode():
        return form_id == 'eures-beta'
    return True


def form_env_suffix(form_id: str | None) -> str:
    """Build the environment-variable suffix for a form id."""
    return str(form_id or '').strip().replace('-', '_').upper()


def _resolve_form_path(form_id: str, raw_path: str) -> str | None:
    """
    Resolve friendly form URLs.
    Supports:
    - explicit file path (existing)
    - "<name>" -> "<name>.html"
    - "<name>/" -> "<name>/index.html" then "<name>.html"
    """
    safe_form_dir = FORMS_DIR / form_id
    candidate = (raw_path or '').strip().lstrip('/')
    if not candidate:
        return None

    candidates: list[str] = []
    candidates.append(candidate)
    if candidate.endswith('/'):
        base = candidate.rstrip('/')
        candidates.append(f"{base}/index.html")
        candidates.append(f"{base}.html")
    else:
        if '.' not in Path(candidate).name:
            candidates.append(f"{candidate}.html")
            candidates.append(f"{candidate}/index.html")

    for c in candidates:
        if (safe_form_dir / c).is_file():
            return c
    return None


def _get_admin_credentials(form_id: str | None = None) -> tuple[str | None, str | None]:
    """Resolve admin credentials, preferring per-form overrides."""
    suffix = form_env_suffix(form_id)
    if suffix:
        username = os.environ.get(f'ADMIN_USERNAME_{suffix}')
        password = os.environ.get(f'ADMIN_PASSWORD_{suffix}')
        if username and password:
            return username, password
    return os.environ.get('ADMIN_USERNAME'), os.environ.get('ADMIN_PASSWORD')


def _get_admin_auth_mode(form_id: str | None = None) -> str:
    """Return the configured admin auth mode for a form."""
    suffix = form_env_suffix(form_id)
    if suffix:
        value = os.environ.get(f'ADMIN_AUTH_MODE_{suffix}')
        if value:
            return value.strip().lower()
    return (os.environ.get('ADMIN_AUTH_MODE') or 'basic').strip().lower() or 'basic'


def _get_admin_allowed_emails(form_id: str | None = None) -> set[str]:
    """Return the normalized allowlist for magic-link admin access."""
    suffix = form_env_suffix(form_id)
    raw = ''
    if suffix:
        raw = os.environ.get(f'ADMIN_ALLOWED_EMAILS_{suffix}', '')
    if not raw:
        raw = os.environ.get('ADMIN_ALLOWED_EMAILS', '')
    return {
        normalize_email(item)
        for item in raw.split(',')
        if normalize_email(item)
    }


def _get_admin_magic_link_ttl_seconds(form_id: str | None = None) -> int:
    """Return the signed-link validity duration."""
    suffix = form_env_suffix(form_id)
    value = ''
    if suffix:
        value = os.environ.get(f'ADMIN_MAGIC_LINK_TTL_SECONDS_{suffix}', '')
    if not value:
        value = os.environ.get('ADMIN_MAGIC_LINK_TTL_SECONDS', '900')
    try:
        return max(60, int(value))
    except (TypeError, ValueError):
        return 900


def _get_admin_magic_link_rate_limit_seconds(form_id: str | None = None) -> int:
    """Return the minimum delay between two link requests."""
    suffix = form_env_suffix(form_id)
    value = ''
    if suffix:
        value = os.environ.get(f'ADMIN_MAGIC_LINK_RATE_LIMIT_SECONDS_{suffix}', '')
    if not value:
        value = os.environ.get('ADMIN_MAGIC_LINK_RATE_LIMIT_SECONDS', '60')
    try:
        return max(0, int(value))
    except (TypeError, ValueError):
        return 60


def _get_admin_session_key(form_id: str | None = None) -> str:
    """Session slot for authenticated admins."""
    return f'admin_auth::{form_id or "global"}'


def _get_admin_magic_link_serializer(form_id: str | None = None) -> URLSafeTimedSerializer:
    """Signer used for admin magic links."""
    return URLSafeTimedSerializer(
        app.secret_key,
        salt=f'admin-magic-link::{form_id or "global"}',
    )


def _is_admin_session_authenticated(form_id: str | None = None) -> bool:
    """Check whether the current browser session is authenticated for the admin."""
    entry = session.get(_get_admin_session_key(form_id))
    if not isinstance(entry, dict):
        return False
    if entry.get('form_id') != form_id:
        return False
    email = normalize_email(entry.get('email'))
    if not email:
        return False
    allowed = _get_admin_allowed_emails(form_id)
    return not allowed or email in allowed


def _set_admin_session_authenticated(form_id: str, email: str) -> None:
    """Persist admin authentication in the Flask session."""
    session[_get_admin_session_key(form_id)] = {
        'form_id': form_id,
        'email': normalize_email(email),
        'authenticated_at': int(time.time()),
    }


def _clear_admin_session(form_id: str | None = None) -> None:
    """Remove admin authentication from the Flask session."""
    session.pop(_get_admin_session_key(form_id), None)


def _render_admin_login_page(form_id: str, *, message: str = '', notice: str = '', error: str = '', email: str = '') -> str:
    """Render the public admin login page."""
    safe_form_id = escape(form_id)
    safe_email = escape(email)
    safe_message = escape(message)
    safe_notice = escape(notice)
    safe_error = escape(error)
    return render_template_string(
        """
<!doctype html>
<html lang="fr">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Connexion admin {{ form_id }}</title>
  <style>
    :root {
      --bg: #f4efe4;
      --panel: #fffdfa;
      --border: #e4d8bf;
      --ink: #1f2f46;
      --muted: #6b6a64;
      --blue: #0f4ea6;
      --blue-deep: #173250;
      --ok: #e9f6ea;
      --ok-border: #b9dfbe;
      --error: #fdecec;
      --error-border: #efb3b3;
    }
    * { box-sizing: border-box; }
    body {
      margin: 0;
      font-family: Georgia, "Times New Roman", serif;
      color: var(--ink);
      background:
        radial-gradient(circle at top right, rgba(244, 205, 93, 0.24), transparent 28%),
        linear-gradient(180deg, #f8f6f1 0%, var(--bg) 100%);
      min-height: 100vh;
      padding: 32px 20px;
    }
    .shell {
      max-width: 720px;
      margin: 0 auto;
    }
    .hero {
      background: linear-gradient(135deg, var(--blue-deep), var(--blue));
      border-radius: 34px 34px 0 0;
      color: white;
      padding: 40px 48px 48px;
    }
    .eyebrow {
      margin: 0 0 16px;
      font-size: 0.95rem;
      letter-spacing: 0.16em;
      text-transform: uppercase;
      opacity: 0.84;
      font-weight: 700;
    }
    h1 {
      margin: 0;
      font-size: clamp(2.2rem, 6vw, 4rem);
      line-height: 0.95;
    }
    .panel {
      background: var(--panel);
      border: 1px solid var(--border);
      border-top: none;
      border-radius: 0 0 34px 34px;
      padding: 42px 48px 48px;
      box-shadow: 0 24px 60px rgba(24, 43, 73, 0.08);
    }
    p {
      margin: 0 0 18px;
      font-size: 1.18rem;
      line-height: 1.6;
    }
    .hint {
      color: var(--muted);
      font-size: 1rem;
    }
    .alert {
      border-radius: 18px;
      padding: 16px 18px;
      margin: 0 0 20px;
      font-size: 1rem;
      line-height: 1.5;
    }
    .alert.notice {
      background: var(--ok);
      border: 1px solid var(--ok-border);
    }
    .alert.error {
      background: var(--error);
      border: 1px solid var(--error-border);
    }
    form {
      margin-top: 28px;
      display: grid;
      gap: 18px;
    }
    label {
      display: grid;
      gap: 8px;
      font-size: 1rem;
      font-weight: 700;
    }
    input[type="email"] {
      width: 100%;
      border: 1px solid #ced5df;
      border-radius: 16px;
      padding: 16px 18px;
      font-size: 1rem;
      color: var(--ink);
      background: white;
    }
    button, .back-link {
      border: none;
      border-radius: 18px;
      padding: 18px 28px;
      font-size: 1rem;
      font-weight: 700;
      text-decoration: none;
      display: inline-flex;
      align-items: center;
      justify-content: center;
    }
    button {
      background: var(--blue);
      color: white;
      cursor: pointer;
      width: fit-content;
    }
    .back-link {
      color: var(--ink);
      background: #eef2f8;
      width: fit-content;
    }
    .actions {
      display: flex;
      gap: 12px;
      flex-wrap: wrap;
      margin-top: 8px;
    }
    @media (max-width: 640px) {
      .hero, .panel {
        padding: 28px 24px 30px;
      }
      p {
        font-size: 1.05rem;
      }
      .actions {
        flex-direction: column;
      }
      button, .back-link {
        width: 100%;
      }
    }
  </style>
</head>
<body>
  <main class="shell">
    <section class="hero">
      <p class="eyebrow">Administration securisee</p>
      <h1>Connexion admin {{ form_id }}</h1>
    </section>
    <section class="panel">
      {% if message %}
      <p>{{ message }}</p>
      {% endif %}
      {% if notice %}
      <div class="alert notice">{{ notice }}</div>
      {% endif %}
      {% if error %}
      <div class="alert error">{{ error }}</div>
      {% endif %}
      <p>Renseignez votre adresse email autorisee pour recevoir un lien de connexion valable 15 minutes.</p>
      <p class="hint">Si vous n'avez pas acces, l'equipe EURES devra ajouter votre adresse a la liste autorisee.</p>
      <form method="post" action="{{ url_for('admin_login', form_id=form_id) }}">
        <label>
          Adresse email
          <input type="email" name="email" value="{{ email }}" required autocomplete="email" inputmode="email">
        </label>
        <div class="actions">
          <button type="submit">Recevoir mon lien de connexion</button>
          <a class="back-link" href="/forms/{{ form_id }}/">Retour au site</a>
        </div>
      </form>
    </section>
  </main>
</body>
</html>
        """,
        form_id=safe_form_id,
        email=safe_email,
        message=safe_message,
        notice=safe_notice,
        error=safe_error,
    )


def _build_admin_magic_link(form_id: str, email: str) -> str:
    """Create the signed admin login URL."""
    token = _get_admin_magic_link_serializer(form_id).dumps({
        'form_id': form_id,
        'email': normalize_email(email),
    })
    return urljoin(
        request.host_url,
        url_for('admin_magic_link_login', form_id=form_id, token=token),
    )


def _send_admin_magic_link_email(form_id: str, email: str) -> None:
    """Send the admin login email through Brevo."""
    api_key = os.environ.get('BREVO_API_KEY', '').strip()
    from_email = os.environ.get('BREVO_FROM_EMAIL', '').strip()
    from_name = os.environ.get('BREVO_FROM_NAME', form_id).strip() or form_id
    if not api_key or not from_email:
        raise RuntimeError('Brevo admin email is not configured.')

    link = _build_admin_magic_link(form_id, email)
    ttl_minutes = max(1, _get_admin_magic_link_ttl_seconds(form_id) // 60)
    subject = f'[{form_id}] Votre lien de connexion admin'
    html_content = f"""
<!doctype html>
<html lang="fr">
  <body style="margin:0;padding:32px;background:#f4efe4;font-family:Georgia,serif;color:#1f2f46;">
    <div style="max-width:720px;margin:0 auto;background:#fffdfa;border:1px solid #e4d8bf;border-radius:34px;overflow:hidden;">
      <div style="padding:40px 48px;background:linear-gradient(135deg,#173250,#0f4ea6);color:#ffffff;">
        <div style="font-size:15px;letter-spacing:0.18em;text-transform:uppercase;font-weight:700;opacity:0.88;">Administration securisee</div>
        <h1 style="margin:18px 0 0;font-size:56px;line-height:0.95;">Connexion admin {escape(form_id)}</h1>
      </div>
      <div style="padding:40px 48px;">
        <p style="font-size:22px;line-height:1.6;margin:0 0 28px;">Bonjour,</p>
        <p style="font-size:22px;line-height:1.6;margin:0 0 28px;">Voici votre lien de connexion pour l'espace d'administration <strong>{escape(form_id)}</strong>.</p>
        <p style="margin:0 0 32px;">
          <a href="{escape(link)}" style="display:inline-block;background:#0f4ea6;color:#ffffff;text-decoration:none;border-radius:18px;padding:18px 28px;font-size:20px;font-weight:700;">Se connecter a l'administration</a>
        </p>
        <p style="font-size:18px;line-height:1.6;margin:0 0 18px;">Ce lien expire dans <strong>{ttl_minutes} minutes</strong>.</p>
        <p style="font-size:18px;line-height:1.6;margin:0 0 28px;color:#6b6a64;">Si vous n'etes pas a l'origine de cette demande, ignorez cet email.</p>
        <p style="font-size:18px;line-height:1.6;margin:0;">Cordialement,<br><br>EURES beta</p>
      </div>
    </div>
  </body>
</html>
    """
    text_content = (
        f"Bonjour,\n\n"
        f"Voici votre lien de connexion pour l'espace d'administration {form_id} :\n\n"
        f"{link}\n\n"
        f"Ce lien expire dans {ttl_minutes} minutes.\n\n"
        f"Si vous n'etes pas a l'origine de cette demande, ignorez cet email.\n"
    )

    response = requests.post(
        'https://api.brevo.com/v3/smtp/email',
        headers={
            'accept': 'application/json',
            'api-key': api_key,
            'content-type': 'application/json',
        },
        json={
            'sender': {'email': from_email, 'name': from_name},
            'to': [{'email': email}],
            'subject': subject,
            'htmlContent': html_content,
            'textContent': text_content,
        },
        timeout=15,
    )
    if response.status_code >= 400:
        raise RuntimeError(f'Brevo email error: {response.status_code} {response.text}')


def _require_basic_admin_auth(form_id: str | None = None):
    """HTTP Basic auth guard for admin endpoints."""
    username, password = _get_admin_credentials(form_id)
    if not username or not password:
        return Response(
            'Admin access not configured. Set ADMIN_USERNAME/ADMIN_PASSWORD or form-specific overrides.',
            503,
            {'Content-Type': 'text/plain; charset=utf-8'},
        )

    auth = request.authorization
    if not auth or auth.username != username or auth.password != password:
        return Response(
            'Authentication required',
            401,
            {'WWW-Authenticate': 'Basic realm="Grist Admin"'},
        )
    return None


def _require_admin_auth(form_id: str | None = None):
    """Protect admin endpoints using the configured auth mode."""
    mode = _get_admin_auth_mode(form_id)
    if mode == 'magic_link':
        if _is_admin_session_authenticated(form_id):
            return None
        basic_denied = _require_basic_admin_auth(form_id)
        if basic_denied is None:
            return None
        if request.path.startswith('/api/'):
            return jsonify({'error': 'Authentication required', 'auth_mode': 'magic_link'}), 401
        return redirect(url_for('admin_login', form_id=form_id, next=request.path))
    return _require_basic_admin_auth(form_id)


def admin_required(fn):
    """Decorator to protect admin pages and APIs."""
    @wraps(fn)
    def wrapper(*args, **kwargs):
        form_id = kwargs.get('form_id')
        if form_id is None and args and isinstance(args[0], str):
            form_id = args[0]
        denied = _require_admin_auth(form_id)
        if denied:
            return denied
        return fn(*args, **kwargs)
    return wrapper


def get_form_config(form_id: str, role: str | None = None) -> dict:
    """Get configuration for a form from environment variables."""
    form_id_upper = form_id.replace('-', '_').upper()
    role_upper = (role or '').strip().replace('-', '_').upper()

    def _env_with_role(base: str) -> str | None:
        if role_upper:
            value = os.environ.get(f'{base}_{form_id_upper}_{role_upper}')
            if value:
                return value
        return os.environ.get(f'{base}_{form_id_upper}')

    doc_id = _env_with_role('GRIST_DOC')
    if not doc_id:
        return None

    return {
        'doc_id': doc_id,
        'table_id': _env_with_role('GRIST_TABLE') or 'Reponses',
        'api_key': _env_with_role('GRIST_API_KEY') or os.environ.get('GRIST_API_KEY'),
    }


def get_eures_stats_config() -> dict | None:
    """Get configuration for the optional EURES monthly stats table."""
    base = get_form_config('eures-beta', 'candidate') or get_form_config('eures-beta')
    if not base:
        return None

    return {
        'doc_id': base['doc_id'],
        'table_id': os.environ.get('GRIST_TABLE_EURES_BETA_STATS', EURES_STATS_TABLE_DEFAULT),
        'api_key': base.get('api_key'),
    }


def get_table_columns(config: dict, headers: dict) -> set[str]:
    """Fetch and cache table column ids to avoid sending unknown fields to Grist."""
    cache_key = (str(config.get('doc_id') or ''), str(config.get('table_id') or ''))
    cached = _TABLE_COLUMNS_CACHE.get(cache_key)
    now = time.time()
    if cached and (now - float(cached.get('loaded_at', 0))) < _TABLE_COLUMNS_CACHE_TTL_SECONDS:
        return set(cached.get('columns', set()))

    url = f"{GRIST_BASE_URL}/api/docs/{config['doc_id']}/tables/{config['table_id']}/columns"
    resp = requests.get(url, headers=headers)
    if resp.status_code != 200:
        raise RuntimeError(f'Failed to read table columns: HTTP {resp.status_code} - {resp.text}')
    payload = resp.json()
    columns = set()
    for col in payload.get('columns', []):
        col_id = (col or {}).get('id')
        if col_id:
            columns.add(str(col_id))
    _TABLE_COLUMNS_CACHE[cache_key] = {
        'columns': columns,
        'loaded_at': now,
    }
    return columns


def ensure_table_columns(config: dict, columns: set[str], headers: dict):
    """Create missing text columns in a Grist table, then refresh cache."""
    if not columns:
        return
    existing = get_table_columns(config, headers)
    missing = sorted(col for col in columns if col not in existing)
    if not missing:
        return

    url = f"{GRIST_BASE_URL}/api/docs/{config['doc_id']}/tables/{config['table_id']}/columns"
    payload = {
        'columns': [
            {
                'id': column_id,
                'fields': {'label': column_id},
                'type': 'Text',
            }
            for column_id in missing
        ]
    }
    resp = requests.post(url, json=payload, headers=headers)
    if resp.status_code != 200:
        raise RuntimeError(f'Failed to create columns in {config["table_id"]}: HTTP {resp.status_code} - {resp.text}')

    cache_key = (str(config.get('doc_id') or ''), str(config.get('table_id') or ''))
    _TABLE_COLUMNS_CACHE.pop(cache_key, None)


def _as_bool(value) -> bool:
    """Normalize truthy values from Grist fields."""
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    txt = str(value).strip().lower()
    return txt in {'1', 'true', 'vrai', 'oui', 'yes'}


def _has_value(value) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    return str(value).strip() != ''


def _safe_json(value, fallback):
    if isinstance(value, str):
        txt = value.strip()
        if not txt:
            return fallback
        try:
            return json.loads(txt)
        except Exception:
            return fallback
    return value if value is not None else fallback


def _details_payload(fields: dict) -> dict:
    """Return parsed prestations_details_json payload as dict."""
    payload = _safe_json((fields or {}).get('prestations_details_json'), {})
    return payload if isinstance(payload, dict) else {}


def _details_get(payload: dict, path: tuple[str, ...], fallback=None):
    """Safely navigate a nested dict payload."""
    cur = payload
    for key in path:
        if not isinstance(cur, dict):
            return fallback
        cur = cur.get(key)
    return cur if cur is not None else fallback


def _has_any_checked(data: dict) -> bool:
    if not isinstance(data, dict):
        return False
    for v in data.values():
        if isinstance(v, dict):
            if _has_any_checked(v):
                return True
        elif _as_bool(v):
            return True
    return False


def _has_any_number_data(data) -> bool:
    if isinstance(data, dict):
        return any(_has_any_number_data(v) for v in data.values())
    if isinstance(data, list):
        return any(_has_any_number_data(v) for v in data)
    if data is None:
        return False
    txt = str(data).strip()
    if txt == '':
        return False
    try:
        return float(txt) >= 0
    except Exception:
        return False


def _format_conditional_display_name(raw_name: str) -> str:
    """Short admin-friendly label for conditional steps."""
    raw = str(raw_name or '').strip()
    if not raw:
        return ''
    parts = [part.strip() for part in raw.split(' - ') if part.strip()]
    if not parts:
        return raw
    if parts[0] == 'Directes ORP CDAPH':
        return 'Prestations directes: ' + ' > '.join(parts[1:]) if len(parts) > 1 else 'Prestations directes'
    if parts[0] == 'Directes hors ORP CDAPH':
        return 'Prestations directes sans ORP CDAPH: ' + ' > '.join(parts[1:]) if len(parts) > 1 else 'Prestations directes sans ORP CDAPH'
    if parts[0] == 'Indirectes':
        return 'Prestations indirectes: ' + ' > '.join(parts[1:]) if len(parts) > 1 else 'Prestations indirectes'
    return raw


def _split_export_path(path: str) -> list[str]:
    """Split flattened paths while keeping list indexes with their segment."""
    return [part for part in str(path or '').split('.') if part]


def _array_index(segment: str) -> tuple[str, int | None]:
    if '[' not in segment or not segment.endswith(']'):
        return segment, None
    name, _, raw_index = segment.partition('[')
    try:
        return name, int(raw_index[:-1]) - 1
    except Exception:
        return name, None


def humanize_export_field(key: str) -> str:
    """Label top-level Grist/export fields with user-facing wording where known."""
    raw = str(key or '')
    if raw in FIELD_LABELS:
        return FIELD_LABELS[raw]
    if raw.startswith('q') and '_' in raw:
        return raw.replace('_', ' ').upper()
    return raw.replace('_', ' ')


def humanize_export_path(path: str) -> str:
    """Turn flattened JSON paths into labels close to the form wording."""
    raw = str(path or '').strip()
    if not raw:
        return ''
    parts = _split_export_path(raw)
    if not parts:
        return humanize_export_field(raw)

    if parts[0] == 'coh' and len(parts) >= 2:
        block_key, index = _array_index(parts[1])
        title, labels = COH_LABELS.get(block_key, (humanize_export_field(block_key), []))
        if index is not None and 0 <= index < len(labels):
            return f'{title} - {labels[index]}'
        return title

    first_name, first_index = _array_index(parts[0])

    if first_name == 'geoRows' and len(parts) >= 2:
        _, index = _array_index(parts[0])
        suffix = EXPORT_KEY_LABELS.get(parts[1], humanize_export_field(parts[1]))
        prefix = f'Géographie ligne {index + 1}' if index is not None else 'Géographie'
        return f'{prefix} - {suffix}'

    if first_name == 'handicapMatrix' and len(parts) >= 2:
        handicap = HANDICAP_TYPES[first_index] if first_index is not None and 0 <= first_index < len(HANDICAP_TYPES) else 'Type de handicap'
        col = EXPORT_KEY_LABELS.get(parts[1], humanize_export_field(parts[1]))
        return f'Type de handicap principal et associé - {handicap} - {col}'

    if parts[0] in {'directAvecOrp', 'directSansOrp', 'indirect', 'pecFileActive', 'preconisationsBloc', 'sortiesBloc', 'suspensionsBloc', 'orienteursBloc', 'formationsSelection'}:
        group_labels = {
            'directAvecOrp': 'Prestations directes avec ORP CDAPH',
            'directSansOrp': 'Prestations directes sans ORP CDAPH',
            'indirect': 'Prestations indirectes',
            'pecFileActive': 'File active PEC',
            'preconisationsBloc': 'Préconisations lors de la sortie définitive',
            'sortiesBloc': 'Sortie définitive avant le terme de la prestation',
            'suspensionsBloc': 'Suspensions',
            'orienteursBloc': 'Orienteurs',
            'formationsSelection': 'Sélection des formations',
        }
        cleaned = []
        for part in parts[1:]:
            name, index = _array_index(part)
            if name in {'row', 'rows', 'raisons'}:
                continue
            label = EXPORT_KEY_LABELS.get(name, humanize_export_field(name))
            if index is not None:
                label = f'{label} {index + 1}'
            cleaned.append(label)
        return ' - '.join([group_labels[parts[0]], *cleaned])

    if len(parts) == 1:
        name, index = _array_index(parts[0])
        label = EXPORT_KEY_LABELS.get(name, humanize_export_field(name))
        return f'{label} {index + 1}' if index is not None else label

    return ' - '.join(EXPORT_KEY_LABELS.get(_array_index(part)[0], humanize_export_field(_array_index(part)[0])) for part in parts)


def should_skip_readable_export_path(path: str) -> bool:
    """Hide internal navigation/state fields that are not useful in a user-readable export."""
    first = _array_index(_split_export_path(path)[0])[0] if _split_export_path(path) else ''
    return first in {
        'visitedBlocks',
        'currentConditionalId',
        'focusedConditionalStepId',
        'focusedConditionalBlockKey',
        'autoCollapsePrimaryAfterPrestations',
    }


def _xlsx_safe_value(value):
    """Return a spreadsheet-friendly scalar without raw JSON objects."""
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'Oui' if value else 'Non'
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _flatten_for_xlsx(value, prefix=''):
    """Flatten nested JSON as readable path/value rows."""
    rows = []
    if isinstance(value, dict):
        if not value:
            rows.append((prefix, ''))
        for key, child in value.items():
            path = f'{prefix}.{key}' if prefix else str(key)
            rows.extend(_flatten_for_xlsx(child, path))
    elif isinstance(value, list):
        if not value:
            rows.append((prefix, ''))
        for idx, child in enumerate(value, start=1):
            path = f'{prefix}[{idx}]' if prefix else f'[{idx}]'
            rows.extend(_flatten_for_xlsx(child, path))
    else:
        rows.append((prefix, _xlsx_safe_value(value)))
    return rows


def _append_table(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append([_xlsx_safe_value(cell) for cell in row])
    for cell in ws[1]:
        cell.style = 'Headline 4'
    ws.freeze_panes = 'A2'
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in column_cells)
        width = min(max(max_len + 2, 12), 80)
        ws.column_dimensions[column_cells[0].column_letter].width = width


def _new_sheet(workbook, title):
    safe = ''.join(ch for ch in str(title or 'Feuille') if ch not in r'[]:*?/\\')[:31] or 'Feuille'
    if safe in workbook.sheetnames:
        base = safe[:27]
        i = 2
        while f'{base} {i}' in workbook.sheetnames:
            i += 1
        safe = f'{base} {i}'
    return workbook.create_sheet(safe)


def build_readable_xlsx(payload: dict) -> BytesIO:
    """Build a human-readable workbook from the current form payload, leaving Grist JSON intact."""
    from openpyxl import Workbook

    fields = payload.get('fields') if isinstance(payload, dict) else {}
    fields = fields if isinstance(fields, dict) else {}

    wb = Workbook()
    wb.remove(wb.active)

    summary_rows = [
        ('Date export', payload.get('exported_at', '')),
        ('UUID', payload.get('uuid', fields.get('uuid', ''))),
        ('Formulaire', payload.get('form_id', '')),
    ]
    for key in sorted(fields):
        if key in JSON_EXPORT_COLUMNS:
            continue
        summary_rows.append((humanize_export_field(key), fields.get(key)))
    _append_table(_new_sheet(wb, 'Synthese'), ['Champ', 'Valeur'], summary_rows)

    metiers = _safe_json(fields.get('metiers_json'), [])
    if isinstance(metiers, list):
        rows = [
            (
                item.get('metier', ''),
                item.get('mode', ''),
                item.get('etp', ''),
                item.get('etpCdi', ''),
                item.get('etpCdd', ''),
            )
            for item in metiers
            if isinstance(item, dict)
        ]
        _append_table(_new_sheet(wb, 'Metiers'), ['Métier', 'Mode', 'ETP', 'ETP CDI', 'ETP CDD'], rows)

    formations = _safe_json(fields.get('catalogue_formations_json'), [])
    if isinstance(formations, list):
        rows = [
            (
                item.get('nom', ''),
                item.get('niveau', ''),
                item.get('secteur', ''),
                item.get('nb', ''),
            )
            for item in formations
            if isinstance(item, dict)
        ]
        _append_table(_new_sheet(wb, 'Formations'), ['Formation', 'Niveau', 'Secteur', 'Nombre'], rows)

    prestations = _safe_json(fields.get('prestations_json'), {})
    prestation_rows = []
    if isinstance(prestations, dict):
        for section_key, section in prestations.items():
            if not isinstance(section, dict):
                continue
            label = _format_conditional_display_name(section_key)
            prestation_rows.append((label, humanize_export_path('done'), section.get('done', '')))
            for path, value in _flatten_for_xlsx(section):
                if path in {'done'} or should_skip_readable_export_path(path):
                    continue
                prestation_rows.append((label, humanize_export_path(path), value))
    _append_table(_new_sheet(wb, 'Prestations'), ['Section', 'Champ', 'Valeur'], prestation_rows)

    details = _details_payload(fields)
    wizard = details.get(WIZARD_STATE_KEY) if isinstance(details, dict) else {}
    runtime = wizard.get('runtime') if isinstance(wizard, dict) else {}
    conditional_defs = runtime.get('conditionalDefs') if isinstance(runtime, dict) else []
    conditional_state = runtime.get('conditionalState') if isinstance(runtime, dict) else {}
    conditional_rows = []
    if isinstance(conditional_defs, list) and isinstance(conditional_state, dict):
        for item in conditional_defs:
            if not isinstance(item, dict):
                continue
            cond_id = str(item.get('id') or '')
            name = _format_conditional_display_name(item.get('name') or cond_id)
            state = conditional_state.get(cond_id, {})
            if not isinstance(state, dict):
                continue
            conditional_rows.append((name, humanize_export_path('done'), state.get('done', '')))
            for path, value in _flatten_for_xlsx(state):
                if path == 'done' or should_skip_readable_export_path(path):
                    continue
                conditional_rows.append((name, humanize_export_path(path), value))
    _append_table(_new_sheet(wb, 'Etapes conditionnelles'), ['Étape', 'Champ', 'Valeur'], conditional_rows)

    selection_rows = []
    for column in ['prestations_orp_json', 'prestations_indirectes_json', 'finess_json']:
        parsed = _safe_json(fields.get(column), {})
        for path, value in _flatten_for_xlsx(parsed, column):
            selection_rows.append((humanize_export_path(path), value))
    _append_table(_new_sheet(wb, 'Selections'), ['Champ', 'Valeur'], selection_rows)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def compute_quick_step_progress(fields: dict) -> dict:
    """
    Fast progress estimator by major sections.
    Returns counts + list of done/remaining step labels.
    """
    finess_values = extract_finess_values(fields)
    identification_done = (
        _has_value(fields.get('es_nom'))
        and _has_value(fields.get('validateur_email'))
        and _has_value(fields.get('es_departement'))
        and len(finess_values) > 0
    )

    selected_dispositifs = [
        _as_bool(fields.get('check_esrp')),
        _as_bool(fields.get('check_espo')),
        _as_bool(fields.get('check_ueros')),
        _as_bool(fields.get('check_deac')),
    ]
    rh_done = any(selected_dispositifs) and _has_value(fields.get('metiers_json'))

    details_json = _details_payload(fields)
    prestations_orp_legacy = _safe_json(fields.get('prestations_orp_json'), {})
    prestations_orp_canonical = _details_get(details_json, ('prestations', 'selection', 'orp'), {})
    prestations_orp = {}
    if isinstance(prestations_orp_canonical, dict):
        prestations_orp.update(prestations_orp_canonical)
    if isinstance(prestations_orp_legacy, dict):
        prestations_orp.update(prestations_orp_legacy)
    vos_prestations_done = _has_value(fields.get('pec')) or _has_any_checked(prestations_orp)

    contexte_done = (
        _has_value(fields.get('q32_implantation'))
        and _has_value(fields.get('q33_transports'))
        and _has_value(fields.get('q53_afpa'))
    )

    steps = [
        ('Identification', identification_done),
        ('Volet RH', rh_done),
        ('Vos prestations', vos_prestations_done),
        ('Contexte écologique', contexte_done),
    ]

    # Conditional quick estimator.
    cond_expected_labels = []
    cond_done_labels = []
    wizard_runtime = _details_get(details_json, (WIZARD_STATE_KEY, 'runtime'), {})

    if isinstance(wizard_runtime, dict):
        runtime_defs = wizard_runtime.get('conditionalDefs')
        runtime_state = wizard_runtime.get('conditionalState')
        if isinstance(runtime_defs, list) and isinstance(runtime_state, dict):
            for item in runtime_defs:
                if not isinstance(item, dict):
                    continue
                cond_id = str(item.get('id') or '').strip()
                cond_name = _format_conditional_display_name(str(item.get('name') or '').strip())
                if not cond_id or not cond_name:
                    continue
                cond_expected_labels.append(cond_name)
                state = runtime_state.get(cond_id)
                if isinstance(state, dict) and (_as_bool(state.get('done')) or _as_bool(state.get('__completed'))):
                    cond_done_labels.append(cond_name)

    if not cond_expected_labels:
        if _as_bool(fields.get('check_esrp')):
            cond_expected_labels.append('Prestations totales ESRP')
        if _as_bool(fields.get('check_espo')):
            cond_expected_labels.append('Prestations totales ESPO')
        if _as_bool(fields.get('check_ueros')):
            cond_expected_labels.append('Prestations totales UEROS')
        if _as_bool(fields.get('check_deac')):
            cond_expected_labels.append('Prestations totales DEAc')

        prestations_json = _safe_json(fields.get('prestations_json'), {})
        prestations_json_canonical = _details_get(details_json, ('prestations', 'conditional', 'state_by_key'), {})
        if (not isinstance(prestations_json, dict) or not prestations_json) and isinstance(prestations_json_canonical, dict):
            prestations_json = prestations_json_canonical
        if isinstance(prestations_json, dict):
            if _has_any_number_data(prestations_json.get('esrp')):
                cond_done_labels.append('Prestations totales ESRP')
            if _has_any_number_data(prestations_json.get('espo')):
                cond_done_labels.append('Prestations totales ESPO')
            if _has_any_number_data(prestations_json.get('ueros')):
                cond_done_labels.append('Prestations totales UEROS')
            if _has_any_number_data(prestations_json.get('deac')):
                cond_done_labels.append('Prestations totales DEAc')

        orp_map = {
            'orp_pec': 'Prestation ORP: PEC',
            'orp_autre_eval': "Prestation ORP: Autre dispositif d'évaluation",
            'orp_orientation_espo': 'Prestation ORP: Orientation ESPO',
            'orp_parcours_sociopro': 'Prestation ORP: Parcours socio-professionnel',
            'orp_autre_parcours': 'Prestation ORP: Autre type de parcours',
            'orp_parcours_qualif': 'Prestation ORP: Parcours certifiant',
            'orp_remise_niveau': 'Prestation ORP: Remise à niveau',
            'orp_formation_pro': 'Prestation ORP: Formation professionnalisante',
            'orp_formation_certif': 'Prestation ORP: Formation certifiante',
            'orp_formation_accomp_pro': 'Prestation ORP: Formation accompagnée pro',
            'orp_formation_accomp_certif': 'Prestation ORP: Formation accompagnée certifiante',
        }
        selected_orp_keys = []
        if isinstance(prestations_orp, dict):
            for key, label in orp_map.items():
                if _as_bool(prestations_orp.get(key)):
                    selected_orp_keys.append((key, label))
                    cond_expected_labels.append(label)

        if isinstance(details_json, dict):
            for key, label in selected_orp_keys:
                state = details_json.get(key, {})
                if not isinstance(state, dict):
                    state = {}
                # New canonical location (schema v2)
                if not state:
                    canonical_state = _details_get(details_json, ('prestations', 'conditional', 'state_by_key', key), {})
                    if isinstance(canonical_state, dict):
                        state = canonical_state
                if isinstance(state, dict) and _as_bool(state.get('__completed')):
                    cond_done_labels.append(label)

    done_main = sum(1 for _, ok in steps if ok)
    total_main = len(steps)
    done_main_labels = [label for label, ok in steps if ok]
    remaining_main = [label for label, ok in steps if not ok]

    cond_expected_labels = list(dict.fromkeys(cond_expected_labels))
    cond_done_labels = list(dict.fromkeys(cond_done_labels))
    remaining_cond = [label for label in cond_expected_labels if label not in cond_done_labels]

    return {
        'main': {
            'done': done_main,
            'total': total_main,
            'done_labels': done_main_labels,
            'remaining_labels': remaining_main,
        },
        'conditional': {
            'done': len(cond_done_labels),
            'total': len(cond_expected_labels),
            'done_labels': cond_done_labels,
            'remaining_labels': remaining_cond,
        },
    }


def fetch_all_records(config: dict, headers: dict):
    """Fetch records for a form table (single pass, up to 5000 rows)."""
    url = f"{GRIST_BASE_URL}/api/docs/{config['doc_id']}/tables/{config['table_id']}/records"
    resp = requests.get(url, params={'limit': 5000}, headers=headers)
    if resp.status_code != 200:
        raise RuntimeError(f'Failed to read records: HTTP {resp.status_code} - {resp.text}')
    payload = resp.json()
    return payload.get('records', [])


def _parse_response_json_safe(resp):
    """Return JSON payload when possible, otherwise a readable fallback dict."""
    try:
        return resp.json()
    except Exception:
        body = (resp.text or '').strip()
        return {
            'error': f'Upstream response is not valid JSON (HTTP {resp.status_code}).',
            'raw': body[:500],
        }


def fetch_record_by_field(base_url: str, field_name: str, value: str, headers: dict):
    """Fetch a single record by a unique field from Grist."""
    filter_param = json.dumps({field_name: [value]})
    resp = requests.get(f"{base_url}/records", params={'filter': filter_param}, headers=headers)
    if resp.status_code != 200:
        return None, resp

    payload = _parse_response_json_safe(resp)
    if not isinstance(payload, dict) or not isinstance(payload.get('records'), list):
        raise RuntimeError('Failed to decode Grist record lookup response.')
    return payload['records'][0] if payload['records'] else None, resp


def resolve_record_key(allowed_columns: set[str]) -> str | None:
    """Pick the unique key used to upsert records for a target table."""
    if 'uuid' in allowed_columns:
        return 'uuid'
    if 'id_tally' in allowed_columns:
        return 'id_tally'
    return None


def _grist_write_redirect_response(resp):
    """Return a readable error when an upstream write endpoint redirects."""
    if not 300 <= resp.status_code < 400:
        return None

    source = str(getattr(resp, 'url', '') or '').strip()
    location = str(resp.headers.get('Location') or '').strip()
    source_part = f' Source URL: {source}.' if source else ''
    suffix = f' Redirect target: {location}' if location else ''
    return jsonify({
        'error': (
            f'Grist write endpoint redirected with HTTP {resp.status_code}. '
            f'Check GRIST_BASE_URL in production.{source_part}{suffix}'
        ),
    }), 502


def write_grist_records(method: str, url: str, payload: dict, headers: dict, max_redirects: int = 3):
    """Follow Grist write redirects while preserving the HTTP method and cookies."""
    current_url = url
    with requests.Session() as session:
        for _ in range(max_redirects + 1):
            resp = session.request(
                method,
                current_url,
                json=payload,
                headers=headers,
                allow_redirects=False,
            )
            if not 300 <= resp.status_code < 400:
                return resp

            location = str(resp.headers.get('Location') or '').strip()
            if not location:
                return resp
            current_url = urljoin(current_url, location)

    return resp


def fetch_table_records(doc_id: str, table_id: str, headers: dict, limit: int = 5000):
    """Read records from a Grist table."""
    url = f"{GRIST_BASE_URL}/api/docs/{doc_id}/tables/{table_id}/records"
    resp = requests.get(url, params={'limit': limit}, headers=headers)
    if resp.status_code != 200:
        raise RuntimeError(f'Failed to read {table_id}: HTTP {resp.status_code} - {resp.text}')
    payload = _parse_response_json_safe(resp)
    return payload.get('records', []) if isinstance(payload, dict) else []


def _month_key(value) -> str | None:
    """Extract YYYY-MM from common timestamp/date values."""
    txt = str(value or '').strip()
    if len(txt) >= 7 and txt[4] == '-':
        return txt[:7]
    for fmt in (
        '%Y-%m-%dT%H:%M:%S.%fZ',
        '%Y-%m-%dT%H:%M:%SZ',
        '%Y-%m-%d',
        '%Y/%m/%d',
        '%d/%m/%Y',
    ):
        try:
            return datetime.strptime(txt, fmt).strftime('%Y-%m')
        except Exception:
            continue
    return None


def _split_multi_value(value) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        raw_items = value
    else:
        raw_items = str(value).split('|')
    return [str(item).strip() for item in raw_items if str(item).strip()]


def _counter_to_rows(counter: Counter, minimum_public_count: int = 5) -> list[dict]:
    """Convert a counter to public rows while masking very small categories."""
    visible: list[dict] = []
    hidden_total = 0
    for label, count in sorted(counter.items(), key=lambda item: (-item[1], str(item[0]).lower())):
        if count < minimum_public_count:
            hidden_total += count
        else:
            visible.append({'label': label, 'count': count})
    if hidden_total:
        visible.append({'label': 'Autres', 'count': hidden_total})
    return visible


def _safe_int(value) -> int:
    try:
        return int(float(value or 0))
    except Exception:
        return 0


def build_eures_public_stats() -> dict:
    """Build a public, non-nominative stats payload for EURES beta."""
    candidate_config = get_form_config('eures-beta', 'candidate')
    employer_config = get_form_config('eures-beta', 'employer')
    if not candidate_config or not employer_config:
        raise RuntimeError('EURES beta candidate/employer Grist configuration is incomplete.')

    candidate_headers = {'Accept': 'application/json'}
    if candidate_config.get('api_key'):
        candidate_headers['Authorization'] = f"Bearer {candidate_config['api_key']}"

    employer_headers = {'Accept': 'application/json'}
    if employer_config.get('api_key'):
        employer_headers['Authorization'] = f"Bearer {employer_config['api_key']}"

    candidats = fetch_table_records(candidate_config['doc_id'], EURES_CANDIDATS_TABLE, candidate_headers)
    besoins = fetch_table_records(employer_config['doc_id'], EURES_BESOINS_TABLE, employer_headers)
    matchings = fetch_table_records(candidate_config['doc_id'], EURES_MATCHINGS_TABLE, candidate_headers)

    monthly: dict[str, dict[str, int | str]] = defaultdict(lambda: {
        'mois': '',
        'candidats': 0,
        'besoins_employeurs': 0,
        'candidats_contactes': 0,
        'candidatures_recues': 0,
        'matchings': 0,
        'candidatures_transmises_employeur': 0,
        'embauches': 0,
    })

    candidats_par_pays = Counter()
    besoins_par_pays = Counter()
    secteurs = Counter()
    mobilite_candidats = Counter()
    matchings_par_statut = Counter()

    for rec in candidats:
        fields = rec.get('fields', {}) if isinstance(rec, dict) else {}
        month = _month_key(fields.get('tally_submitted_at'))
        if month:
            monthly[month]['mois'] = month
            monthly[month]['candidats'] += 1
            monthly[month]['candidatures_recues'] += 1
        for label in _split_multi_value(fields.get('pays')):
            candidats_par_pays[label] += 1
        for label in _split_multi_value(fields.get('metier')):
            secteurs[label] += 1
        for label in _split_multi_value(fields.get('mobilite')):
            mobilite_candidats[label] += 1

    for rec in besoins:
        fields = rec.get('fields', {}) if isinstance(rec, dict) else {}
        month = _month_key(fields.get('tally_submitted_at'))
        if month:
            monthly[month]['mois'] = month
            monthly[month]['besoins_employeurs'] += 1
        for label in _split_multi_value(fields.get('pays') or fields.get('pays_normalise')):
            besoins_par_pays[label] += 1
        for label in _split_multi_value(fields.get('poste')):
            secteurs[label] += 1

    for rec in matchings:
        fields = rec.get('fields', {}) if isinstance(rec, dict) else {}
        month = _month_key(fields.get('date_calcul'))
        if month:
            monthly[month]['mois'] = month
            monthly[month]['matchings'] += 1
        status = str(fields.get('statut') or '').strip()
        if status:
            matchings_par_statut[status] += 1

    manual_stats_available = False
    stats_config = get_eures_stats_config()
    if stats_config:
        stats_headers = {'Accept': 'application/json'}
        if stats_config.get('api_key'):
            stats_headers['Authorization'] = f"Bearer {stats_config['api_key']}"
        try:
            manual_rows = fetch_table_records(stats_config['doc_id'], stats_config['table_id'], stats_headers)
            manual_stats_available = True
        except Exception:
            manual_rows = []
        for rec in manual_rows:
            fields = rec.get('fields', {}) if isinstance(rec, dict) else {}
            month = _month_key(fields.get('mois'))
            if not month:
                continue
            monthly[month]['mois'] = month
            monthly[month]['candidats_contactes'] += _safe_int(fields.get('candidats_contactes'))
            monthly[month]['candidatures_transmises_employeur'] += _safe_int(fields.get('candidatures_transmises_employeur'))
            monthly[month]['embauches'] += _safe_int(fields.get('embauches'))

    monthly_rows = [row for _, row in sorted(monthly.items()) if row.get('mois')]

    totals = {
        'candidats': len(candidats),
        'besoins_employeurs': len(besoins),
        'candidats_contactes': sum(int(row['candidats_contactes']) for row in monthly_rows),
        'candidatures_recues': len(candidats),
        'matchings': len(matchings),
        'candidatures_transmises_employeur': sum(int(row['candidatures_transmises_employeur']) for row in monthly_rows),
        'embauches': sum(int(row['embauches']) for row in monthly_rows),
    }

    return {
        'ok': True,
        'generated_at': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
        'manual_stats_table': {
            'configured': manual_stats_available,
            'table_id': stats_config['table_id'] if stats_config else None,
        },
        'totals': totals,
        'monthly': monthly_rows,
        'breakdowns': {
            'candidats_par_pays': _counter_to_rows(candidats_par_pays),
            'besoins_par_pays': _counter_to_rows(besoins_par_pays),
            'secteurs': _counter_to_rows(secteurs),
            'mobilite_candidats': _counter_to_rows(mobilite_candidats),
            'matchings_par_statut': _counter_to_rows(matchings_par_statut),
        },
    }


def fetch_matching_record(doc_id: str, besoin_id: str, candidat_id: str, headers: dict):
    """Read one matching by besoin_id/candidat_id pair."""
    url = f"{GRIST_BASE_URL}/api/docs/{doc_id}/tables/{EURES_MATCHINGS_TABLE}/records"
    filter_param = json.dumps({'besoin_id': [besoin_id], 'candidat_id': [candidat_id]})
    resp = requests.get(url, params={'filter': filter_param}, headers=headers)
    if resp.status_code != 200:
        raise RuntimeError(f'Failed to read Matchings: HTTP {resp.status_code} - {resp.text}')
    payload = _parse_response_json_safe(resp)
    records = payload.get('records', []) if isinstance(payload, dict) else []
    return records[0] if records else None


def upsert_matching_record(doc_id: str, fields: dict, headers: dict):
    """Create or update a matching record from besoin_id/candidat_id."""
    besoin_id = str(fields.get('besoin_id') or '').strip()
    candidat_id = str(fields.get('candidat_id') or '').strip()
    if not besoin_id or not candidat_id:
        raise RuntimeError('Missing besoin_id or candidat_id for matching upsert.')

    table_config = {'doc_id': doc_id, 'table_id': EURES_MATCHINGS_TABLE}
    ensure_table_columns(table_config, set(fields.keys()) & EURES_MATCHING_FIELDS, headers)
    allowed_columns = get_table_columns(table_config, headers)
    filtered_fields = {k: v for k, v in fields.items() if k in allowed_columns}
    existing = fetch_matching_record(doc_id, besoin_id, candidat_id, headers)
    base_url = f"{GRIST_BASE_URL}/api/docs/{doc_id}/tables/{EURES_MATCHINGS_TABLE}/records"

    if existing:
        payload = {'records': [{'id': existing['id'], 'fields': filtered_fields}]}
        resp = write_grist_records('PATCH', base_url, payload, headers)
    else:
        payload = {'records': [{'fields': filtered_fields}]}
        resp = write_grist_records('POST', base_url, payload, headers)

    if resp.status_code != 200:
        raise RuntimeError(f'Failed to write Matchings: HTTP {resp.status_code} - {resp.text}')
    return resp


def eures_normalize_text(value) -> str:
    return str(value or '').strip().lower()


def eures_canonical_sector(value: str) -> str:
    value_l = eures_normalize_text(value)
    for token, sector in EURES_SECTOR_CANONICAL_MAP.items():
        if token in value_l:
            return sector
    return ''


def eures_candidate_sectors(value: str) -> set[str]:
    value_l = eures_normalize_text(value)
    return {sector for token, sector in EURES_SECTOR_CANONICAL_MAP.items() if token in value_l}


def eures_parse_language_requirements(value: str) -> dict[str, int]:
    levels = {
        'pas nécessaire': 0,
        'pas necessaire': 0,
        'not necessary': 0,
        'communication professionelle': 2,
        'communication professionnelle': 2,
        'professional communication': 2,
        'pouvoir communiquer dans la langue locale': 2,
        'très important au quotidien': 3,
        'tres important au quotidien': 3,
        'very important on a daily basis': 3,
    }
    result: dict[str, int] = {}
    for part in str(value or '').split('|'):
        if ':' not in part:
            continue
        lang, raw_level = part.split(':', 1)
        lang_key = eures_normalize_text(lang)
        level_text = eures_normalize_text(raw_level)
        score = next((v for k, v in levels.items() if k in level_text), None)
        if score is not None:
            result[lang_key] = score
    return result


def eures_parse_candidate_languages(value: str) -> dict[str, int]:
    levels = {
        'je ne la pratique pas': 0,
        'i do not speak it': 0,
        "j'utilise un outil de traduction si nécessaire": 0,
        "j'utilise un outil de traduction si necessaire": 0,
        'translation tool': 0,
        'je peux communiquer simplement': 2,
        "je peux échanger à l'oral": 2,
        "je peux echanger a l'oral": 2,
        'simple communication': 2,
        'je peux travailler avec cette langue': 3,
        'langue maternelle': 4,
        'native': 4,
    }
    result: dict[str, int] = {}
    for part in str(value or '').split('|'):
        if ':' not in part:
            continue
        lang, raw_level = part.split(':', 1)
        lang_key = eures_normalize_text(lang)
        level_text = eures_normalize_text(raw_level)
        matched = [v for k, v in levels.items() if k in level_text]
        if matched:
            result[lang_key] = max(matched)
    return result


def eures_availability_rank(value: str) -> int | None:
    value_l = eures_normalize_text(value)
    ordered = [
        ('dès que possible', 0),
        ('des que possible', 0),
        ('dans les prochains jours', 1),
        ('dans les prochaines semaines', 2),
        ('dans 1 à 3 mois', 3),
        ('dans 1 a 3 mois', 3),
        ('je ne sais pas encore', 4),
    ]
    for token, rank in ordered:
        if token in value_l:
            return rank
    return None


def eures_to_float(value):
    if value in (None, '', 0, '0'):
        return None
    try:
        return float(value)
    except Exception:
        return None


def eures_score_sector_fit(expected: str, actual: str) -> tuple[int, str]:
    expected_sector = eures_canonical_sector(expected)
    candidate_sectors = eures_candidate_sectors(actual)
    if not expected_sector:
        return 0, 'metier: secteur employeur absent'
    if not candidate_sectors:
        return 0, 'metier: secteur candidat absent'
    if expected_sector in candidate_sectors:
        if len(candidate_sectors) == 1:
            return 30, f'metier: secteur exact ({expected_sector})'
        return 20, f'metier: secteur present ({expected_sector})'
    return 0, f'metier: secteur non aligne ({expected_sector})'


def eures_score_languages(expected: str, actual: str) -> tuple[int, str]:
    expected_map = eures_parse_language_requirements(expected)
    actual_map = eures_parse_candidate_languages(actual)
    scored = {lang: lvl for lang, lvl in expected_map.items() if lvl > 0}
    if not scored or not actual_map:
        return 0, 'langues: information manquante'
    matched = []
    strong = 0
    partial = 0
    for lang, required_level in scored.items():
        candidate_level = actual_map.get(lang)
        if candidate_level is None:
            continue
        if candidate_level >= required_level:
            strong += 1
            matched.append(lang)
        elif candidate_level >= 2 and candidate_level + 1 >= required_level:
            partial += 1
            matched.append(f'{lang} (partiel)')
    ratio = (strong + 0.5 * partial) / max(len(scored), 1)
    points = round(25 * ratio)
    if points > 0:
        return points, 'langues: ' + ', '.join(matched)
    return 0, 'langues: aucune correspondance'


def eures_score_location(besoin_pays: str, candidat_pays: str, mobilite: str) -> tuple[int, str]:
    besoin_pays_l = eures_normalize_text(besoin_pays)
    candidat_pays_l = eures_normalize_text(candidat_pays)
    mobilite_l = eures_normalize_text(mobilite)
    if besoin_pays_l and candidat_pays_l and (besoin_pays_l == candidat_pays_l or besoin_pays_l in candidat_pays_l):
        return 15, 'pays/mobilite: meme pays'
    if besoin_pays_l and mobilite_l and besoin_pays_l in mobilite_l:
        return 15, 'pays/mobilite: pays souhaite explicite'
    if mobilite_l and any(token in mobilite_l for token in ('transfrontali', 'expatriation', 'pays souhaités', 'pays souhaites')):
        return 10, 'pays/mobilite: mobilite declaree'
    return 0, 'pays/mobilite: contrainte geographique'


def eures_score_availability(date_debut: str, disponibilite: str) -> tuple[int, str]:
    disponibilite_l = str(disponibilite or '').strip()
    if not disponibilite_l:
        return 0, 'disponibilite: information manquante'
    besoin_rank = eures_availability_rank(date_debut)
    candidat_rank = eures_availability_rank(disponibilite)
    if besoin_rank is not None and candidat_rank is not None:
        if candidat_rank <= besoin_rank:
            return 15, 'disponibilite: date compatible'
        if candidat_rank == besoin_rank + 1:
            return 9, 'disponibilite: leger decalage'
        return 0, 'disponibilite: delai trop long'
    if eures_normalize_text(date_debut) and eures_normalize_text(date_debut) in eures_normalize_text(disponibilite):
        return 15, 'disponibilite: date compatible'
    return 7, 'disponibilite: verification manuelle recommandee'


def eures_candidate_salary_expectation(secteur: str, fields: dict):
    salary_fields = EURES_CANDIDAT_SECTOR_SALARY_FIELDS.get(secteur)
    if not salary_fields:
        return None
    type_field, min_field = salary_fields
    salary_type = str(fields.get(type_field) or '').strip()
    salary_min = eures_to_float(fields.get(min_field))
    if salary_min is None:
        return None
    return salary_type, salary_min


def eures_employer_salary_offer(secteur: str, fields: dict):
    salary_fields = EURES_EMPLOYEUR_SECTOR_SALARY_FIELDS.get(secteur)
    if not salary_fields:
        return None
    type_field, min_field, max_field = salary_fields
    salary_type = str(fields.get(type_field) or '').strip()
    salary_min = eures_to_float(fields.get(min_field))
    salary_max = eures_to_float(fields.get(max_field))
    if salary_min is None and salary_max is None:
        return None
    return salary_type, salary_min, salary_max


def eures_score_salary(secteur: str, candidat_fields: dict, besoin_fields: dict) -> tuple[int, str]:
    if not secteur:
        return 0, 'salaire: secteur indetermine'
    candidat_salary = eures_candidate_salary_expectation(secteur, candidat_fields)
    employeur_salary = eures_employer_salary_offer(secteur, besoin_fields)
    if candidat_salary is None or employeur_salary is None:
        return 6, 'salaire: donnees absentes, verification manuelle'
    _, candidat_min = candidat_salary
    _, employeur_min, employeur_max = employeur_salary
    if employeur_max and candidat_min <= employeur_max:
        if employeur_min and candidat_min < employeur_min:
            return 12, 'salaire: attente sous la fourchette proposee'
        return 15, 'salaire: compatible'
    if employeur_min and candidat_min <= employeur_min * 1.1:
        return 8, 'salaire: ecart limite'
    return 0, 'salaire: attente superieure au salaire propose'


def eures_matching_status(score: int) -> str:
    if score >= 75:
        return 'auto_envoyable'
    if score >= 55:
        return 'a_valider'
    return 'a_ecarter'


def compute_eures_matching(besoin_fields: dict, candidat_fields: dict) -> dict:
    """Compute a matching payload for one candidate/employer pair."""
    besoin_country = besoin_fields.get('pays_normalise') or besoin_fields.get('pays') or ''
    candidat_country = candidat_fields.get('pays_normalise') or candidat_fields.get('pays') or ''
    secteur = eures_canonical_sector(str(besoin_fields.get('poste') or ''))

    score_metier, raison_metier = eures_score_sector_fit(
        str(besoin_fields.get('poste') or ''),
        str(candidat_fields.get('metier') or ''),
    )
    score_langues, raison_langues = eures_score_languages(
        str(besoin_fields.get('langues_requises') or ''),
        str(candidat_fields.get('langues') or ''),
    )
    score_mobilite, raison_mobilite = eures_score_location(
        str(besoin_country),
        str(candidat_country),
        str(candidat_fields.get('mobilite') or ''),
    )
    score_disponibilite, raison_disponibilite = eures_score_availability(
        str(besoin_fields.get('date_debut') or ''),
        str(candidat_fields.get('disponibilite') or ''),
    )
    score_salaire, raison_salaire = eures_score_salary(secteur, candidat_fields, besoin_fields)

    reasons = []
    weaknesses = []
    for points, text in [
        (score_metier, raison_metier),
        (score_langues, raison_langues),
        (score_mobilite, raison_mobilite),
        (score_disponibilite, raison_disponibilite),
        (score_salaire, raison_salaire),
    ]:
        (reasons if points else weaknesses).append(text)

    score = score_metier + score_langues + score_mobilite + score_disponibilite + score_salaire
    return {
        'score': score,
        'score_metier': score_metier,
        'score_langues': score_langues,
        'score_mobilite': score_mobilite,
        'score_disponibilite': score_disponibilite,
        'score_salaire': score_salaire,
        'statut': eures_matching_status(score),
        'raisons': ' | '.join(reasons),
        'points_faibles': ' | '.join(weaknesses),
        'date_calcul': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
    }


def run_eures_matching_for_saved_record(form_id: str, role: str, saved_record: dict, config: dict, headers: dict):
    """Compute matchings for one newly saved EURES beta record and write them to Matchings."""
    if form_id != 'eures-beta':
        return {'processed': False, 'reason': 'unsupported_form'}
    if role not in {'candidate', 'employer'}:
        return {'processed': False, 'reason': 'unsupported_role'}

    all_candidats = fetch_table_records(config['doc_id'], EURES_CANDIDATS_TABLE, headers)
    all_besoins = fetch_table_records(config['doc_id'], EURES_BESOINS_TABLE, headers)
    saved_fields = saved_record.get('fields', {}) if isinstance(saved_record.get('fields'), dict) else {}

    writes = 0
    if role == 'candidate':
        for besoin in all_besoins:
            besoin_fields = besoin.get('fields', {}) if isinstance(besoin.get('fields'), dict) else {}
            candidat_id = str(saved_fields.get('id_tally') or saved_fields.get('uuid') or '')
            besoin_id = str(besoin_fields.get('id_tally') or besoin_fields.get('uuid') or '')
            if not candidat_id or not besoin_id:
                continue
            matching = compute_eures_matching(besoin_fields, saved_fields)
            payload = {'besoin_id': besoin_id, 'candidat_id': candidat_id, **matching}
            upsert_matching_record(config['doc_id'], payload, headers)
            writes += 1
    else:
        for candidat in all_candidats:
            candidat_fields = candidat.get('fields', {}) if isinstance(candidat.get('fields'), dict) else {}
            candidat_id = str(candidat_fields.get('id_tally') or candidat_fields.get('uuid') or '')
            besoin_id = str(saved_fields.get('id_tally') or saved_fields.get('uuid') or '')
            if not candidat_id or not besoin_id:
                continue
            matching = compute_eures_matching(saved_fields, candidat_fields)
            payload = {'besoin_id': besoin_id, 'candidat_id': candidat_id, **matching}
            upsert_matching_record(config['doc_id'], payload, headers)
            writes += 1

    return {'processed': True, 'writes': writes, 'role': role}


def normalize_finess(value) -> str:
    """Normalize FINESS value for duplicate checks."""
    raw = str(value or '').strip()
    digits = ''.join(ch for ch in raw if ch.isdigit())
    if digits and digits == raw:
        return digits.zfill(9) if len(digits) == 8 else digits
    return raw.upper()

def normalize_email(value) -> str:
    """Normalize email value for lookup."""
    return str(value or '').strip().lower()


def extract_finess_values(fields: dict) -> set:
    """Extract FINESS values from standard and JSON fields."""
    values = set()
    if not isinstance(fields, dict):
        return values

    main = normalize_finess(fields.get('finess_main'))
    if main:
        values.add(main)

    raw_json = fields.get('finess_json')
    parsed = []
    if isinstance(raw_json, str) and raw_json.strip():
        try:
            parsed = json.loads(raw_json)
        except Exception:
            parsed = []
    elif isinstance(raw_json, list):
        parsed = raw_json

    if isinstance(parsed, list):
        for item in parsed:
            v = normalize_finess(item)
            if v:
                values.add(v)

    return values


def find_duplicate_finess(config: dict, current_uuid: str, finess_values: set, headers: dict):
    """Find FINESS values already used by another record in the same table."""
    if not finess_values:
        return set()

    records = fetch_all_records(config, headers)
    duplicates = set()
    cur = normalize_finess(current_uuid)

    for rec in records:
        fields = rec.get('fields', {}) if isinstance(rec, dict) else {}
        rec_uuid = normalize_finess(fields.get('uuid'))
        if cur and rec_uuid == cur:
            continue
        rec_finess = extract_finess_values(fields)
        duplicates.update(finess_values.intersection(rec_finess))

    return duplicates


@app.route('/api/forms/<form_id>/record', methods=['GET'])
def get_record(form_id: str):
    """Fetch a record by UUID or table-specific identifier."""
    if not is_form_enabled(form_id):
        return jsonify({'error': f'Unknown form: {form_id}'}), 404
    config = get_form_config(form_id, request.args.get('flow_role'))
    if not config:
        return jsonify({'error': f'Unknown form: {form_id}'}), 404

    uuid = request.args.get('uuid')
    if not uuid:
        return jsonify({'error': 'UUID parameter required'}), 400

    # Build Grist API URL with filter
    filter_param = f'{{"uuid":["{uuid}"]}}'
    url = f"{GRIST_BASE_URL}/api/docs/{config['doc_id']}/tables/{config['table_id']}/records"

    headers = {'Accept': 'application/json'}
    # API key optional for read if doc is public
    if config['api_key']:
        headers['Authorization'] = f"Bearer {config['api_key']}"

    try:
        allowed_columns = get_table_columns(config, headers)
        record_key = resolve_record_key(allowed_columns)
        if not record_key:
            return jsonify({'error': 'Target table has no supported record key (uuid or id_tally).'}), 500
        resp = requests.get(
            url,
            params={'filter': json.dumps({record_key: [uuid]})},
            headers=headers,
        )
        return jsonify(resp.json()), resp.status_code
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/forms/<form_id>/record', methods=['POST'])
def save_record(form_id: str):
    """Create or update a record."""
    if not is_form_enabled(form_id):
        return jsonify({'error': f'Unknown form: {form_id}'}), 404
    data = request.get_json()
    if not data or 'fields' not in data:
        return jsonify({'error': 'Invalid request body'}), 400

    fields = data['fields']
    if not isinstance(fields, dict):
        return jsonify({'error': 'Invalid fields payload'}), 400
    config = get_form_config(form_id, fields.get('flow_role'))
    if not config:
        return jsonify({'error': f'Unknown form: {form_id}'}), 404

    if not config['api_key']:
        return jsonify({'error': 'API key not configured for this form'}), 500
    uuid = fields.get('uuid') or fields.get('id_tally')
    if not uuid:
        return jsonify({'error': 'UUID or id_tally required in fields'}), 400

    base_url = f"{GRIST_BASE_URL}/api/docs/{config['doc_id']}/tables/{config['table_id']}"
    headers = {
        'Authorization': f"Bearer {config['api_key']}",
        'Content-Type': 'application/json',
        'Accept': 'application/json',
    }

    # Keep only fields that exist in target table (prod/local may differ).
    try:
        allowed_columns = get_table_columns(config, headers)
        filtered_fields = {k: v for k, v in fields.items() if str(k) in allowed_columns}
        record_key = resolve_record_key(allowed_columns)
    except Exception as e:
        return jsonify({'error': f'Failed to fetch Grist table columns: {e}'}), 500

    if not record_key:
        return jsonify({'error': "Table is missing a supported unique key column ('uuid' or 'id_tally')"}), 500

    if record_key not in filtered_fields:
        filtered_fields[record_key] = uuid

    # Check if record exists
    try:
        existing_record, check_resp = fetch_record_by_field(base_url, record_key, filtered_fields[record_key], headers)
        if check_resp.status_code != 200:
            return jsonify(_parse_response_json_safe(check_resp)), check_resp.status_code
        record_id = existing_record.get('id') if isinstance(existing_record, dict) else None
    except Exception as e:
        return jsonify({'error': f'Failed to check existing record: {e}'}), 500

    # Enforce uniqueness of FINESS across records (excluding current UUID)
    try:
        incoming_finess = extract_finess_values(filtered_fields)
        duplicates = find_duplicate_finess(config, uuid, incoming_finess, headers)
        if duplicates:
            duplicates_sorted = sorted(duplicates)
            return jsonify({
                'error': 'Un ou plusieurs numéros FINESS sont déjà utilisés par un autre questionnaire.',
                'duplicates': duplicates_sorted,
            }), 409
    except Exception as e:
        return jsonify({'error': f'Failed to validate FINESS uniqueness: {e}'}), 500

    # Create or update, then re-read the UUID so callers never get a false success.
    try:
        action = 'updated' if record_id else 'created'
        if record_id:
            # Update existing
            payload = {'records': [{'id': record_id, 'fields': filtered_fields}]}
            resp = write_grist_records('PATCH', f"{base_url}/records", payload, headers)
        else:
            # Create new
            payload = {'records': [{'fields': filtered_fields}]}
            resp = write_grist_records('POST', f"{base_url}/records", payload, headers)

        redirect_error = _grist_write_redirect_response(resp)
        if redirect_error:
            return redirect_error

        if resp.status_code != 200:
            return jsonify(_parse_response_json_safe(resp)), resp.status_code

        saved_record, verify_resp = fetch_record_by_field(base_url, record_key, filtered_fields[record_key], headers)
        if verify_resp.status_code != 200:
            return jsonify({
                'error': f'Grist write returned success, but verification failed with HTTP {verify_resp.status_code}.',
            }), 502
        if not isinstance(saved_record, dict):
            return jsonify({
                'error': 'Grist write returned success, but the questionnaire was not found after saving.',
            }), 502

        saved_fields = saved_record.get('fields', {}) if isinstance(saved_record.get('fields'), dict) else {}
        if str(saved_fields.get(record_key) or '') != str(filtered_fields[record_key]):
            return jsonify({
                'error': f'Grist write returned success, but the saved questionnaire {record_key} did not match.',
            }), 502

        matching_result = None
        if form_id == 'eures-beta':
            matching_result = run_eures_matching_for_saved_record(
                form_id=form_id,
                role=str(fields.get('flow_role') or ''),
                saved_record=saved_record,
                config=config,
                headers=headers,
            )

        return jsonify({
            'ok': True,
            'action': action,
            'uuid': uuid,
            'record_key': record_key,
            'record_id': saved_record.get('id'),
            'matching': matching_result,
        }), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/forms/<form_id>/export-readable-xlsx', methods=['POST'])
def export_readable_xlsx(form_id: str):
    """Generate a human-readable Excel export from a form payload without changing Grist storage."""
    if not is_form_enabled(form_id):
        return jsonify({'error': f'Unknown form: {form_id}'}), 404
    data = request.get_json()
    if not isinstance(data, dict):
        return jsonify({'error': 'Invalid request body'}), 400
    try:
        xlsx = build_readable_xlsx(data)
    except Exception as e:
        return jsonify({'error': f'Failed to build readable Excel export: {e}'}), 500

    uuid = str(data.get('uuid') or (data.get('fields') or {}).get('uuid') or 'saisie').strip() or 'saisie'
    filename = f'fagerh_saisie_lisible_{uuid}.xlsx'
    return send_file(
        xlsx,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


@app.route('/api/forms/<form_id>/check-finess', methods=['POST'])
def check_finess(form_id: str):
    """Check whether FINESS values already exist in another questionnaire."""
    if not is_form_enabled(form_id):
        return jsonify({'error': f'Unknown form: {form_id}'}), 404
    config = get_form_config(form_id)
    if not config:
        return jsonify({'error': f'Unknown form: {form_id}'}), 404

    data = request.get_json() or {}
    uuid = normalize_finess(data.get('uuid'))
    raw_finess = data.get('finess') or []

    if not isinstance(raw_finess, list):
        return jsonify({'error': 'Invalid request body: finess must be a list'}), 400

    finess_values = {normalize_finess(v) for v in raw_finess if normalize_finess(v)}
    headers = {'Accept': 'application/json'}
    if config['api_key']:
        headers['Authorization'] = f"Bearer {config['api_key']}"

    try:
        duplicates = sorted(find_duplicate_finess(config, uuid, finess_values, headers))
        return jsonify({
            'ok': True,
            'duplicates': duplicates,
            'has_duplicates': len(duplicates) > 0,
        }), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/forms/<form_id>/recover-by-email', methods=['POST'])
def recover_by_email(form_id: str):
    """Recover a questionnaire UUID from validation email (+ optional FINESS)."""
    if not is_form_enabled(form_id):
        return jsonify({'error': f'Unknown form: {form_id}'}), 404
    config = get_form_config(form_id)
    if not config:
        return jsonify({'error': f'Unknown form: {form_id}'}), 404

    data = request.get_json() or {}
    email = normalize_email(data.get('email'))
    finess = normalize_finess(data.get('finess'))
    if not email:
        return jsonify({'error': 'Email required'}), 400
    if finess and not finess.isdigit():
        return jsonify({'error': 'FINESS invalide'}), 400

    url = f"{GRIST_BASE_URL}/api/docs/{config['doc_id']}/tables/{config['table_id']}/records"
    headers = {'Accept': 'application/json'}
    if config['api_key']:
        headers['Authorization'] = f"Bearer {config['api_key']}"

    try:
        # 1) Fast path: exact filter on the email column.
        filter_param = json.dumps({"validateur_email": [email]})
        resp = requests.get(url, params={'filter': filter_param, 'limit': 5000}, headers=headers)
        if resp.status_code != 200:
            return jsonify(resp.json()), resp.status_code
        payload = resp.json()
        records = payload.get('records', [])

        matches = []
        for rec in records:
            fields = rec.get('fields', {}) if isinstance(rec, dict) else {}
            if not fields.get('uuid'):
                continue
            if finess:
                rec_finess = extract_finess_values(fields)
                if finess not in rec_finess:
                    continue
            matches.append(rec)

        # 2) Fallback: scan and compare case-insensitively (older rows / casing differences).
        if not matches:
            scan_resp = requests.get(url, params={'limit': 5000}, headers=headers)
            if scan_resp.status_code != 200:
                return jsonify(scan_resp.json()), scan_resp.status_code
            scan_payload = scan_resp.json()
            for rec in scan_payload.get('records', []):
                fields = rec.get('fields', {}) if isinstance(rec, dict) else {}
                if normalize_email(fields.get('validateur_email')) != email:
                    continue
                if not fields.get('uuid'):
                    continue
                if finess:
                    rec_finess = extract_finess_values(fields)
                    if finess not in rec_finess:
                        continue
                matches.append(rec)

        if not matches:
            if finess:
                return jsonify({'error': 'Aucun questionnaire trouvé pour ce couple email + FINESS.'}), 404
            return jsonify({'error': 'Aucun questionnaire trouvé pour cet email.'}), 404

        chosen = max(matches, key=lambda r: int(r.get('id', 0) or 0))
        chosen_fields = chosen.get('fields', {}) if isinstance(chosen, dict) else {}
        return jsonify({
            'ok': True,
            'uuid': chosen_fields.get('uuid'),
            'count': len(matches),
            'match_on': 'email+finess' if finess else 'email',
        }), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/forms/<form_id>/admin/overview', methods=['GET'])
@admin_required
def admin_overview(form_id: str):
    """Admin dashboard data: counts + questionnaire list."""
    if not is_form_enabled(form_id):
        return jsonify({'error': f'Unknown form: {form_id}'}), 404
    config = get_form_config(form_id)
    if not config:
        return jsonify({'error': f'Unknown form: {form_id}'}), 404

    headers = {'Accept': 'application/json'}
    if config['api_key']:
        headers['Authorization'] = f"Bearer {config['api_key']}"

    search = str(request.args.get('search', '') or '').strip().lower()
    status = str(request.args.get('status', 'all') or 'all').strip().lower()

    try:
        records = fetch_all_records(config, headers)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

    rows = []
    for rec in records:
        fields = rec.get('fields', {}) if isinstance(rec, dict) else {}
        if not fields:
            continue
        row = {
            'record_id': rec.get('id'),
            'uuid': fields.get('uuid', ''),
            'es_nom': fields.get('es_nom', ''),
            'departement': fields.get('es_departement', ''),
            'finess_main': fields.get('finess_main', ''),
            'validateur_nom': fields.get('validateur_nom', ''),
            'validateur_prenom': fields.get('validateur_prenom', ''),
            'validateur_email': fields.get('validateur_email', ''),
            'saisie_terminee': _as_bool(fields.get('saisie_terminee')),
        }
        row['quick_progress'] = compute_quick_step_progress(fields)
        rows.append(row)

    total = len(rows)
    termines = sum(1 for r in rows if r['saisie_terminee'])
    en_cours = total - termines

    # Newest first by record id.
    rows.sort(key=lambda r: int(r['record_id'] or 0), reverse=True)

    if status in {'en_cours', 'termines'}:
        want_done = status == 'termines'
        rows = [r for r in rows if r['saisie_terminee'] is want_done]

    if search:
        def _matches(r):
            haystack = ' '.join([
                str(r.get('es_nom', '')),
                str(r.get('departement', '')),
                str(r.get('finess_main', '')),
                str(r.get('validateur_nom', '')),
                str(r.get('validateur_prenom', '')),
                str(r.get('validateur_email', '')),
                str(r.get('uuid', '')),
            ]).lower()
            return search in haystack
        rows = [r for r in rows if _matches(r)]

    return jsonify({
        'ok': True,
        'form_id': form_id,
        'stats': {
            'total_questionnaires': total,
            'en_cours': en_cours,
            'termines': termines,
        },
        'rows': rows,
    }), 200


@app.route('/api/forms/<form_id>/public-stats', methods=['GET'])
def public_stats(form_id: str):
    """Public aggregated stats page data."""
    if not is_form_enabled(form_id) or form_id != 'eures-beta':
        return jsonify({'error': f'Unknown public stats form: {form_id}'}), 404

    try:
        return jsonify(build_eures_public_stats()), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/admin/<form_id>/login', methods=['GET', 'POST'])
def admin_login(form_id: str):
    """Public entry point for magic-link admin authentication."""
    if not is_form_enabled(form_id):
        return jsonify({'error': 'File not found'}), 404

    if _get_admin_auth_mode(form_id) != 'magic_link':
        return redirect(url_for('serve_admin', form_id=form_id))

    if request.method == 'GET':
        if _is_admin_session_authenticated(form_id):
            return redirect(url_for('serve_admin', form_id=form_id))
        notice = request.args.get('notice', '')
        error = request.args.get('error', '')
        return _render_admin_login_page(form_id, notice=notice, error=error)

    email = normalize_email(request.form.get('email'))
    allowed_emails = _get_admin_allowed_emails(form_id)
    if not email:
        return _render_admin_login_page(form_id, error='Renseignez une adresse email valide.'), 400
    if allowed_emails and email not in allowed_emails:
        return _render_admin_login_page(
            form_id,
            error="Cette adresse email n'est pas autorisee pour l'administration.",
            email=email,
        ), 403

    now_ts = int(time.time())
    rate_limit = _get_admin_magic_link_rate_limit_seconds(form_id)
    rate_limit_key = f'admin_magic_last_sent::{form_id}::{email}'
    last_sent = int(session.get(rate_limit_key, 0) or 0)
    wait_seconds = rate_limit - (now_ts - last_sent)
    if rate_limit and wait_seconds > 0:
        return _render_admin_login_page(
            form_id,
            error=f"Un lien vient deja d'etre demande. Attendez encore {wait_seconds} seconde(s).",
            email=email,
        ), 429

    try:
        _send_admin_magic_link_email(form_id, email)
    except Exception as exc:
        return _render_admin_login_page(
            form_id,
            error=f"Impossible d'envoyer le lien de connexion : {exc}",
            email=email,
        ), 503

    session[rate_limit_key] = now_ts
    return _render_admin_login_page(
        form_id,
        notice=f"Un lien de connexion a ete envoye a {email}.",
        email=email,
    )


@app.route('/admin/<form_id>/magic')
def admin_magic_link_login(form_id: str):
    """Consume a signed magic link and authenticate the admin session."""
    if not is_form_enabled(form_id):
        return jsonify({'error': 'File not found'}), 404

    token = request.args.get('token', '').strip()
    if not token:
        return redirect(url_for('admin_login', form_id=form_id, error='Lien de connexion invalide.'))

    try:
        payload = _get_admin_magic_link_serializer(form_id).loads(
            token,
            max_age=_get_admin_magic_link_ttl_seconds(form_id),
        )
    except SignatureExpired:
        return redirect(url_for('admin_login', form_id=form_id, error='Ce lien de connexion a expire.'))
    except BadSignature:
        return redirect(url_for('admin_login', form_id=form_id, error='Lien de connexion invalide.'))

    email = normalize_email((payload or {}).get('email'))
    if (payload or {}).get('form_id') != form_id or not email:
        return redirect(url_for('admin_login', form_id=form_id, error='Lien de connexion invalide.'))

    allowed_emails = _get_admin_allowed_emails(form_id)
    if allowed_emails and email not in allowed_emails:
        return redirect(url_for('admin_login', form_id=form_id, error="Cette adresse email n'est plus autorisee."))

    _set_admin_session_authenticated(form_id, email)
    return redirect(url_for('serve_admin', form_id=form_id))


@app.route('/admin/<form_id>/logout', methods=['POST', 'GET'])
def admin_logout(form_id: str):
    """Clear admin session for magic-link authenticated users."""
    _clear_admin_session(form_id)
    if _get_admin_auth_mode(form_id) == 'magic_link':
        return redirect(url_for('admin_login', form_id=form_id, notice='Vous etes deconnecte.'))
    return redirect(url_for('serve_admin', form_id=form_id))


@app.route('/health')
def health():
    """Health check endpoint."""
    return jsonify({'status': 'ok'})


@app.route('/')
def index():
    if is_eures_beta_only_mode():
        return redirect('/forms/eures-beta/')
    return redirect('/forms/fagerh/')


# Static file serving for forms and frontend
@app.route('/forms/<form_id>/')
@app.route('/forms/<form_id>')
def serve_form(form_id: str):
    """Serve form HTML."""
    if not is_form_enabled(form_id):
        return jsonify({'error': 'File not found'}), 404
    return send_from_directory(FORMS_DIR / form_id, 'index.html')


@app.route('/forms/fagerh/questions-pdf')
def serve_fagerh_questions_pdf():
    """Serve the reference PDF containing FAGERH questions."""
    if is_eures_beta_only_mode():
        return jsonify({'error': 'File not found'}), 404
    return send_from_directory(
        DOCS_DIR,
        'fagerh_questions_completes.pdf',
        as_attachment=True,
        download_name='fagerh_questions_fagerh.pdf',
    )


@app.route('/forms/<form_id>/<path:filename>')
def serve_form_file(form_id: str, filename: str):
    """Serve extra files from a form folder (e.g. UI prototypes)."""
    if not is_form_enabled(form_id):
        return jsonify({'error': 'File not found'}), 404
    resolved = _resolve_form_path(form_id, filename)
    if not resolved:
        return jsonify({'error': 'File not found'}), 404
    return send_from_directory(FORMS_DIR / form_id, resolved)


@app.route('/admin/<form_id>/')
@app.route('/admin/<form_id>')
@admin_required
def serve_admin(form_id: str):
    """Serve admin dashboard HTML for a form."""
    if not is_form_enabled(form_id):
        return jsonify({'error': 'File not found'}), 404
    return send_from_directory(FORMS_DIR / form_id, 'admin.html')


@app.route('/assets/<path:filename>')
def serve_assets(filename: str):
    """Serve static assets (JS, CSS)."""
    return send_from_directory(ASSETS_DIR, filename)

if fagerh_suivi_app is not None and not is_eures_beta_only_mode():
    app.wsgi_app = DispatcherMiddleware(
        app.wsgi_app,
        {
            '/forms/fagerh/saisie-quotidienne': fagerh_suivi_app,
        },
    )


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5005, debug=False)
